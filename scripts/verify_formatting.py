import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

doc_formatted = docx.Document(r"C:\Users\Dell_PC\Downloads\MultiCDSS.docx")
doc_backup = docx.Document(r"C:\Users\Dell_PC\Downloads\MultiCDSS.original.backup.docx")

print("--- 1. Verification of Text & Structure ---")
assert len(doc_formatted.paragraphs) == len(doc_backup.paragraphs), "Paragraph count mismatch!"
for idx, (p1, p2) in enumerate(zip(doc_formatted.paragraphs, doc_backup.paragraphs)):
    assert p1.text == p2.text, f"Paragraph {idx} text differs!"
print(f"Checked {len(doc_formatted.paragraphs)} paragraphs: 100% IDENTICAL")

assert len(doc_formatted.tables) == len(doc_backup.tables), "Table count mismatch!"
for t_i, (t1, t2) in enumerate(zip(doc_formatted.tables, doc_backup.tables)):
    assert len(t1.rows) == len(t2.rows), f"Table {t_i} rows differ!"
    for r_i, (row1, row2) in enumerate(zip(t1.rows, t2.rows)):
        assert len(row1.cells) == len(row2.cells), f"Table {t_i} row {r_i} cols differ!"
        for c_i, (c1, c2) in enumerate(zip(row1.cells, row2.cells)):
            assert c1.text == c2.text, f"Table {t_i} cell ({r_i},{c_i}) differs!"
print(f"Checked {len(doc_formatted.tables)} tables: 100% IDENTICAL")

print("\n--- 2. Verification of Images / Figures ---")
orig_blips = [i for i, p in enumerate(doc_backup.paragraphs) if 'blip' in p._p.xml]
curr_blips = [i for i, p in enumerate(doc_formatted.paragraphs) if 'blip' in p._p.xml]
assert orig_blips == curr_blips, "Image blip paragraphs differ!"
print(f"Image blips preserved at paragraphs: {curr_blips}")

print("\n--- 3. Verification of Headings and Colors ---")
DARK_BLUE = "1F4E79"
BLUE = "2F75B5"
BLACK = "000000"
WHITE = "FFFFFF"

# Check Main Title (P1)
p1_colors = {str(r.font.color.rgb) for r in doc_formatted.paragraphs[1].runs if r.font.color}
p1_bold = all(r.bold for r in doc_formatted.paragraphs[1].runs)
print(f"P1 (Main Title): colors={p1_colors}, bold={p1_bold}")
assert p1_colors == {DARK_BLUE}
assert p1_bold == True

# Check Main Headings
for p_idx in [11, 18, 22, 36, 108, 111]:
    p = doc_formatted.paragraphs[p_idx]
    colors = {str(r.font.color.rgb) for r in p.runs if r.font.color}
    bold = all(r.bold for r in p.runs if r.text.strip())
    print(f"P{p_idx} ({p.text.strip()[:20]}): colors={colors}, bold={bold}")
    assert colors == {DARK_BLUE}
    assert bold == True

# Check Subsection Headings
for p_idx in [37, 43, 44, 49, 55, 65, 70]:
    p = doc_formatted.paragraphs[p_idx]
    colors = {str(r.font.color.rgb) for r in p.runs if r.font.color}
    bold = all(r.bold for r in p.runs if r.text.strip())
    print(f"P{p_idx} ({p.text.strip()[:20]}): colors={colors}, bold={bold}")
    assert colors == {BLUE}
    assert bold == True

# Check Table Headers (Dark Blue shading + White text)
print("\n--- 4. Verification of Table Headers and Shading ---")
for t_idx in range(2, len(doc_formatted.tables)):
    t = doc_formatted.tables[t_idx]
    for c in t.rows[0].cells:
        tcPr = c._tc.get_or_add_tcPr()
        shd_xml = [elem for elem in tcPr if elem.tag.endswith('shd')][0].attrib
        fill = shd_xml.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill', '')
        assert fill.upper() == DARK_BLUE, f"Table {t_idx} header cell shading is {fill}, expected {DARK_BLUE}"
        for p in c.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    assert str(r.font.color.rgb) == WHITE, f"Table {t_idx} header text color {r.font.color.rgb} != WHITE"
    print(f"Table {t_idx} header verified: Dark Blue shading #{DARK_BLUE}, White text #{WHITE}")

print("\n--- 5. Verification of Body Text Color ---")
non_black_body_runs = []
for i, p in enumerate(doc_formatted.paragraphs):
    if i in [1, 11, 18, 22, 36, 108, 111, 37, 43, 44, 49, 55, 65, 70, 25, 90, 96, 104, 40, 72, 77, 84, 6, 7, 10, 28, 29, 32, 33, 34, 35]:
        continue
    for r in p.runs:
        if r.font.color and r.font.color.rgb and str(r.font.color.rgb) != BLACK:
            non_black_body_runs.append((i, r.text, str(r.font.color.rgb)))

assert len(non_black_body_runs) == 0, f"Found non-black body runs: {non_black_body_runs}"
print(f"Verified all general body paragraphs: 100% BLACK (#{BLACK})")

print("\n============================================================")
print("ALL VERIFICATIONS COMPLETED SUCCESSFULLY WITH ZERO ERRORS!")
print("============================================================")
